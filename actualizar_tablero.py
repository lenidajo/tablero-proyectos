"""Actualiza docs/index.html a partir de CronoEstadistica!.xlsx y publica el cambio en git.

Flujo:
  1. Lee el Excel fuente (hoja "Crono", datos desde la fila 5).
  2. Convierte cada fila a un dict: dep, concepto, tipo, f_agg, f_ini, f_fin, pct, obs.
  3. Inserta ese JSON en template.html (reemplaza __DATA_JSON__) y guarda docs/index.html.
  4. git add / commit / push.

Si el Excel no existe en la ruta esperada, el script se detiene sin tocar docs/index.html.
"""

import datetime
import json
import re
import subprocess
import sys
from pathlib import Path

import openpyxl

EXCEL_PATH = Path(r"C:\Users\daniel.clavijo\Desktop\Informacion Estadistica\CronoEstadistica!.xlsx")
SHEET_NAME = "Crono"
FIRST_DATA_ROW = 5

SCRIPT_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = SCRIPT_DIR / "template.html"
OUTPUT_PATH = SCRIPT_DIR / "docs" / "index.html"

# columnas B..I en la hoja "Crono"
COL_DEP, COL_CONCEPTO, COL_TIPO, COL_F_AGG, COL_F_INI, COL_F_FIN, COL_PCT, COL_OBS = range(2, 10)


def fail(msg):
    print(f"ERROR: {msg}", file=sys.stderr)
    sys.exit(1)


def normalize_date(value):
    if value is None:
        return None
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        if re.match(r"^\d{4}-\d{2}-\d{2}$", s):
            return s
        m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", s)
        if m:
            d, mo, y = m.groups()
            return f"{y}-{int(mo):02d}-{int(d):02d}"
        m = re.match(r"^(\d{1,2})-(\d{1,2})-(\d{2})$", s)
        if m:
            d, mo, y = m.groups()
            return f"{2000 + int(y)}-{int(mo):02d}-{int(d):02d}"
        return s
    return str(value)


def clean_str(value):
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def read_projects():
    if not EXCEL_PATH.exists():
        fail(
            f'no se encontró el Excel fuente en "{EXCEL_PATH}". '
            "No se publican datos viejos ni inventados; corrige la ruta o el archivo y vuelve a correr el script."
        )

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        fail(f'la hoja "{SHEET_NAME}" no existe en el Excel. Hojas disponibles: {wb.sheetnames}')
    ws = wb[SHEET_NAME]

    rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        dep = ws.cell(r, COL_DEP).value
        concepto = ws.cell(r, COL_CONCEPTO).value
        tipo = ws.cell(r, COL_TIPO).value
        f_agg = ws.cell(r, COL_F_AGG).value
        f_ini = ws.cell(r, COL_F_INI).value
        f_fin = ws.cell(r, COL_F_FIN).value
        pct = ws.cell(r, COL_PCT).value
        obs = ws.cell(r, COL_OBS).value

        if all(v is None for v in [dep, concepto, tipo, f_agg, f_ini, f_fin, pct, obs]):
            continue

        rows.append(
            {
                "dep": clean_str(dep),
                "concepto": clean_str(concepto),
                "tipo": clean_str(tipo),
                "f_agg": normalize_date(f_agg),
                "f_ini": normalize_date(f_ini),
                "f_fin": normalize_date(f_fin),
                "pct": pct if isinstance(pct, (int, float)) else 0,
                "obs": clean_str(obs),
            }
        )

    if not rows:
        fail("el Excel no tiene filas de datos desde la fila 5. No se publica un tablero vacío.")

    return rows


def build_html(projects):
    if not TEMPLATE_PATH.exists():
        fail(f"no se encontró template.html en {TEMPLATE_PATH}")

    template = TEMPLATE_PATH.read_text(encoding="utf-8")
    if "__DATA_JSON__" not in template:
        fail("template.html no contiene el placeholder __DATA_JSON__")

    data_json = json.dumps(projects, ensure_ascii=False)
    html = template.replace("__DATA_JSON__", data_json)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(html, encoding="utf-8")
    return html


def run_git(*args, check=True):
    return subprocess.run(
        ["git", *args],
        cwd=SCRIPT_DIR,
        capture_output=True,
        text=True,
        check=check,
    )


def git_publish(commit_message):
    result = run_git("status", "--porcelain")
    if not result.stdout.strip():
        print("Sin cambios respecto al último commit; no se publica nada nuevo.")
        return

    run_git("add", "-A")

    diff = run_git("diff", "--cached", "--quiet", check=False)
    if diff.returncode == 0:
        print("Sin cambios en el índice tras 'git add'; no se hace commit.")
        return

    commit = run_git("commit", "-m", commit_message, check=False)
    if commit.returncode != 0:
        fail(f"git commit falló:\n{commit.stdout}\n{commit.stderr}")
    print(commit.stdout.strip())

    push = run_git("push", check=False)
    if push.returncode != 0:
        print(
            "AVISO: el commit se creó localmente pero 'git push' falló "
            "(probablemente falta configurar el remoto o autenticarse en GitHub).\n"
            f"{push.stdout}\n{push.stderr}",
            file=sys.stderr,
        )
        return
    print(push.stdout.strip())
    print(push.stderr.strip())


def main():
    projects = read_projects()

    total = len(projects)
    avg_pct = round(100 * sum(p["pct"] for p in projects) / total)
    urgentes = sum(1 for p in projects if p["tipo"] == "Urgente")

    build_html(projects)

    today = datetime.date.today().isoformat()
    commit_message = f"{today}: {total} proyectos, {avg_pct}% avance promedio, {urgentes} urgentes"

    git_publish(commit_message)

    print(f"OK — {total} proyectos, {avg_pct}% avance promedio, {urgentes} urgentes.")
    print(f"docs/index.html actualizado en {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
